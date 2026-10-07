import sys, json, docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
sys.stdout.reconfigure(encoding="utf-8")
path=sys.argv[1]; res=json.load(open("results.json",encoding="utf-8"))["results"]
d=docx.Document(path)
t=[t for t in d.tables if len(t.columns)==6 and t.rows[0].cells[4].text.strip()=="ผล TOR"][-1]
body=t.rows[1:]
assert len(body)==len(res),(len(body),len(res))
def shade(cell,color):
    tcPr=cell._tc.get_or_add_tcPr()
    for s in tcPr.findall(qn("w:shd")): tcPr.remove(s)
    s=OxmlElement("w:shd"); s.set(qn("w:val"),"clear"); s.set(qn("w:color"),"auto"); s.set(qn("w:fill"),color); tcPr.append(s)
def settext(cell,text):
    p=cell.paragraphs[0]; runs=p.runs
    runs[0].text=text
    for r in runs[1:]: r.text=""
sub={"not_supported":("ตรง – ไม่ระบุใน DS","F4B183"),"wording_conflict":("ตรง – ถ้อยคำขัดกับ DS","F4B183"),
     "unverifiable":("ตรง – ยืนยันไม่ได้ (ดูภาพ)","F4B183"),"partially_supported":("ตรง – DS ไม่ครบ ควรตรวจซ้ำ","FFEB9C")}
for row,r in zip(body,res):
    c=row.cells
    if r["reference_check"]=="match":
        txt,col=sub.get(r["evidence_support"],(None,None))
        if not txt and r["evidence_support"]=="fully_supported" and r["confidence"]!="high": txt,col="ตรง – ควรตรวจซ้ำ","FFEB9C"
        if txt: settext(c[2],txt); shade(c[2],col)
    if r["reference_check"] in ("mismatch","not_found"):
        shade(c[1],"FFC7CE")
    if r["tor_decision"] in ("compliant","better") and r["confidence"]=="medium":
        settext(c[4],("ผ่าน" if r["tor_decision"]=="compliant" else "ดีกว่า")+" – ควรตรวจซ้ำ")
        if r["reference_check"] not in ("mismatch","not_found"): shade(c[4],"F4B183")
        if not r["key_issue"].startswith("ควรตรวจซ้ำ"): settext(c[5],"ควรตรวจซ้ำ: "+c[5].text)
    print(r["item"],"|",c[1].text,"|",c[2].text,"|",c[3].text[:30],"|",c[4].text,"|",c[5].text[:90])
d.save(path)
