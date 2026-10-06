from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


LABELS = {
    "match": "ตรง", "mismatch": "ไม่ตรง",
    "not_found": "ไม่พบเลขหน้า", "unverifiable": "ยืนยันไม่ได้",
    "not_applicable": "ไม่เกี่ยวข้อง",
}
COLORS = {
    "match": "C6EFCE", "mismatch": "FFC7CE",
    "not_found": "FFC7CE", "unverifiable": "FFEB9C", "not_applicable": "D9EAD3",
}

SUMMARY_SPECS = [
    ("สรุปหน้าอ้างอิง", "reference_check",
     [("match", "ตรง"), ("mismatch", "ไม่ตรง"), ("unverifiable", "ยืนยันไม่ได้"),
      ("not_found", "ไม่พบ"), ("not_applicable", "ไม่เกี่ยวข้อง")]),
    ("สรุป Highlight", "highlight_check",
     [("complete", "ครบ"), ("partial", "บางส่วน"), ("not_found", "ไม่พบ"),
      ("unverifiable", "ยืนยันไม่ได้"), ("not_applicable", "ไม่เกี่ยวข้อง")]),
    ("สรุปการรองรับหลักฐาน", "evidence_support",
     [("fully_supported", "รองรับครบ"), ("partially_supported", "รองรับบางส่วน"),
      ("not_supported", "ไม่รองรับ"), ("wording_conflict", "ถ้อยคำขัดกัน"),
      ("unverifiable", "ยืนยันไม่ได้")]),
    ("สรุปผล TOR", "tor_decision",
     [("compliant", "ผ่าน"), ("better", "ดีกว่า"), ("non_compliant", "ไม่ผ่าน"),
      ("mixed", "ผสม"), ("unverifiable", "ยืนยันไม่ได้"), ("not_applicable", "ไม่เกี่ยวข้อง")]),
    ("สรุปการตรวจ Comply/Better เดิม", "declared_status_check",
     [("match", "ตรง"), ("mismatch", "ไม่ตรง"), ("not_selected", "ไม่ได้เลือก"),
      ("ambiguous", "กำกวม"), ("not_applicable", "ไม่เกี่ยวข้อง")]),
]

SUMMARY_COLORS = {
    "match": "C6EFCE", "complete": "C6EFCE", "fully_supported": "C6EFCE", "compliant": "C6EFCE",
    "better": "A9D18E", "mismatch": "FFC7CE", "not_supported": "FFC7CE", "non_compliant": "FFC7CE",
    "partial": "FFEB9C", "partially_supported": "FFEB9C", "wording_conflict": "FFEB9C",
    "mixed": "FFEB9C", "unverifiable": "FFEB9C", "ambiguous": "FFEB9C",
    "not_found": "FFC7CE", "not_selected": "E7E6E6", "not_applicable": "D9EAD3",
}

HIGHLIGHT_LABELS = {
    "complete": "ครบ", "partial": "บางส่วน", "not_found": "ไม่พบ",
    "unverifiable": "ยืนยันไม่ได้", "not_applicable": "ไม่เกี่ยวข้อง",
}
EVIDENCE_LABELS = {
    "fully_supported": "รองรับครบ", "partially_supported": "รองรับบางส่วน",
    "not_supported": "ไม่รองรับ", "wording_conflict": "ถ้อยคำขัดกัน",
    "unverifiable": "ยืนยันไม่ได้", "not_applicable": "ไม่เกี่ยวข้อง",
}
TOR_LABELS = {
    "compliant": "ผ่าน", "better": "ดีกว่า", "non_compliant": "ไม่ผ่าน",
    "mixed": "ผสม", "unverifiable": "ยืนยันไม่ได้", "not_applicable": "ไม่เกี่ยวข้อง",
}


def fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_width(cell, cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(Cm(cm).twips)))
    tc_w.set(qn("w:type"), "dxa")


def put_text(cell, value: object, *, bold: bool = False, size: float = 8) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("" if value is None else str(value))
    run.bold = bold
    run.font.name = "TH Sarabun New"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
    run.font.size = Pt(size)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def add_summary(doc, results: list[dict], title: str, field: str,
                statuses: list[tuple[str, str]]) -> None:
    present = [(status, label) for status, label in statuses if any(item.get(field) == status for item in results)]
    if not present:
        return
    paragraph = doc.add_paragraph()
    put_run = paragraph.add_run(title)
    put_run.bold = True
    put_run.font.name = "TH Sarabun New"
    put_run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
    put_run.font.size = Pt(10)
    counts = Counter(item.get(field) for item in results)
    table = doc.add_table(rows=2, cols=len(present))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, (status, label) in enumerate(present):
        put_text(table.cell(0, index), label, bold=True)
        fill(table.cell(0, index), SUMMARY_COLORS.get(status, "E7E6E6"))
        put_text(table.cell(1, index), counts.get(status, 0), bold=True, size=10)
        table.cell(1, index).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER


def require_fields(item: dict, fields: list[str], index: int, context: str) -> None:
    missing = [field for field in fields if field not in item or item[field] in (None, "")]
    if missing:
        raise ValueError(f"{context} requires {', '.join(missing)} at result {index}")


def validate(data: dict) -> list[dict]:
    if not isinstance(data, dict):
        raise ValueError("results.json root must be an object")
    mode = data.get("mode")
    if mode not in {"standard", "full_audit", "custom"}:
        raise ValueError("mode must be standard, full_audit, or custom")
    options = data.get("options", [])
    checks = data.get("checks", [])
    if not isinstance(options, list) or not all(isinstance(value, str) for value in options):
        raise ValueError("options must be a list of strings")
    if not isinstance(checks, list) or not all(isinstance(value, str) for value in checks):
        raise ValueError("checks must be a list of strings")

    active = set(options) | set(checks)
    if mode in {"standard", "full_audit"}:
        active.update({"reference", "heading_title", "product_identity", "content_relevance"})
    if mode == "full_audit":
        active.update({"item_label", "highlight"})
    if "tor_decision" in active and "evidence_support" not in active:
        raise ValueError("tor_decision requires evidence_support")

    results = data.get("results")
    if not isinstance(results, list) or not results:
        raise ValueError("results.json must contain a non-empty results list")
    for index, item in enumerate(results, start=1):
        if not isinstance(item, dict):
            raise ValueError(f"result {index} must be an object")
        if item.get("reference_check") not in LABELS:
            raise ValueError(f"invalid reference_check at result {index}")
        require_fields(item, ["row", "item", "detail"], index, "every result")
        if "reference" in active:
            require_fields(item, ["reference_detail"], index, "reference check")
        if "item_label" in active:
            require_fields(item, ["item_label_check"], index, "item-label check")
        if "highlight" in active:
            require_fields(item, ["highlight_check", "highlight_evidence"], index, "highlight check")
        if "evidence_support" in active:
            require_fields(item, ["evidence_support", "evidence_detail"], index, "evidence-support audit")
        if "tor_decision" in active:
            require_fields(item, [
                "tor_decision", "tor_decision_basis", "verified_value", "tor_threshold",
                "tor_claim_results", "declared_status", "declared_status_check",
            ], index, "TOR decision")
            if not isinstance(item["tor_claim_results"], list):
                raise ValueError(f"tor_claim_results must be a list at result {index}")
    return results


def compact_text(value: object, limit: int = 300) -> str:
    text = " ".join(str(value).split())
    return text if len(text) <= limit else text[:limit - 1].rstrip() + "…"


def clean_issue_text(value: object, limit: int = 180) -> str:
    """Keep human-facing findings; drop extraction metrics and internal status syntax."""
    kept = []
    blocked = ("coverage=", "รองรับคำสำคัญ", "รองรับ 0%", "highlight=", "row_type=",
               "token_coverage", "numeric_coverage", "tor_claim_results")
    for segment in str(value).replace(" | ", "; ").split(";"):
        segment = " ".join(segment.split()).strip(" ,")
        lower = segment.lower()
        if segment and not any(marker in lower for marker in blocked):
            kept.append(segment)
    text = "; ".join(kept)
    for prefix in ("complete_reference:", "mixed_reference:", "wrong_reference:", "missing_reference:"):
        if text.lower().startswith(prefix):
            text = text[len(prefix):].strip()
            break
    return compact_text(text, limit) if text else ""


def compose_audit_status(item: dict) -> str:
    parts = []
    if item.get("highlight_check"):
        status = item["highlight_check"]
        parts.append(f"Highlight: {HIGHLIGHT_LABELS.get(status, status)}")
    if item.get("evidence_support"):
        status = item["evidence_support"]
        parts.append(f"หลักฐาน: {EVIDENCE_LABELS.get(status, status)}")
    return "\n".join(parts) or "—"


def compose_key_issue(item: dict) -> str:
    if item.get("key_issue"):
        return compact_text(item["key_issue"])

    parts = []
    rules = [
        ("reference_check", {"mismatch", "not_found", "unverifiable"}, "reference_detail"),
        ("heading_title_check", {"mismatch", "unverifiable"}, "heading_title_detail"),
        ("highlight_check", {"partial", "not_found", "unverifiable"}, "highlight_evidence"),
        ("evidence_support", {"partially_supported", "not_supported", "wording_conflict", "unverifiable"}, "evidence_detail"),
        ("tor_decision", {"better", "non_compliant", "mixed", "unverifiable"}, "tor_decision_basis"),
        ("declared_status_check", {"mismatch", "ambiguous"}, "declared_status_detail"),
    ]
    for status_field, exceptional, detail_field in rules:
        if item.get(status_field) in exceptional and item.get(detail_field):
            cleaned = clean_issue_text(item[detail_field])
            if cleaned:
                parts.append(cleaned)
    if item.get("candidate_reference"):
        parts.append(f"หน้าแนะนำ: {compact_text(item['candidate_reference'], 80)}")
    if str(item.get("confidence", "")).lower() == "low":
        parts.append("ความมั่นใจต่ำ ควรตรวจภาพต้นฉบับซ้ำ")

    if not parts:
        raw = str(item.get("detail", ""))
        debug_markers = ("row_type=", "token_coverage", "numeric_coverage", "tor_claim_results")
        if raw and not any(marker in raw for marker in debug_markers):
            cleaned = clean_issue_text(raw, 220)
            if cleaned:
                parts.append(cleaned)

    unique = list(dict.fromkeys(part for part in parts if part))
    if unique:
        return compact_text("; ".join(unique))
    if item.get("reference_check") == "match" and item.get("evidence_support") == "fully_supported":
        return "หน้าอ้างอิงและหลักฐานรองรับข้อกำหนดครบ"
    if item.get("reference_check") == "match":
        return "หน้าอ้างอิงถูกต้อง ไม่พบประเด็นสำคัญ"
    return "ไม่พบรายละเอียดสำคัญเพิ่มเติม"


def build(input_docx: Path, results_json: Path, output_docx: Path) -> None:
    if input_docx.resolve() == output_docx.resolve():
        raise ValueError("input and output must be different files")
    data = json.loads(results_json.read_text(encoding="utf-8"))
    results = validate(data)
    doc = Document(input_docx)

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    # Force width > height instead of blindly swapping: the source document may
    # already be landscape (common for wide TOR/SOC comparison tables), and a
    # blind swap would flip it back to portrait, causing the results table to
    # overflow the page. max()/min() is idempotent regardless of source orientation.
    wide, narrow = max(section.page_width, section.page_height), min(section.page_width, section.page_height)
    section.page_width, section.page_height = wide, narrow
    section.top_margin = section.bottom_margin = Cm(1.2)
    section.left_margin = section.right_margin = Cm(1.2)
    section.header.is_linked_to_previous = False
    for table in list(section.header.tables):
        table._element.getparent().remove(table._element)
    for paragraph in list(section.header.paragraphs):
        paragraph._element.getparent().remove(paragraph._element)
    section.header.add_paragraph("")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(data.get("title", "ผลการตรวจสอบ TOR/SOC เทียบหลักฐาน"))
    run.bold = True
    run.font.name = "TH Sarabun New"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
    run.font.size = Pt(18)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(f"วันที่ตรวจ {data.get('audit_date', '-')} | ผลรายข้อแนบท้ายเอกสารต้นฉบับ")
    run.font.name = "TH Sarabun New"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
    run.font.size = Pt(10)

    for summary_title, field, statuses in SUMMARY_SPECS:
        add_summary(doc, results, summary_title, field, statuses)

    doc.add_paragraph()
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    # Column widths are ratios, not fixed cm — they get scaled to whatever the
    # actual usable width of this section is (page width minus margins). This
    # keeps the table from overflowing regardless of the source document's
    # page size/orientation (landscape A4, portrait A4 forced to landscape, etc).
    width_ratios = [1.6, 4.0, 2.8, 4.0, 2.6, 10.0]
    usable_width_cm = (section.page_width - section.left_margin - section.right_margin) / 360000
    scale = usable_width_cm / sum(width_ratios)
    widths = [ratio * scale for ratio in width_ratios]
    headers = ["ข้อ", "หน้าอ้างอิง", "ผลอ้างอิง", "Highlight/หลักฐาน", "ผล TOR", "ประเด็นหลัก"]
    for index, (label, cell_width) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[index]
        set_width(cell, cell_width)
        fill(cell, "D9EAF7")
        put_text(cell, label, bold=True)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    repeat_header(table.rows[0])

    for item in results:
        status = item["reference_check"]
        tor_status = item.get("tor_decision", "not_applicable")
        values = [item["item"], item.get("reference", "—") or "—", LABELS[status],
                  compose_audit_status(item), TOR_LABELS.get(tor_status, tor_status),
                  compose_key_issue(item)]
        cells = table.add_row().cells
        for cell, value, cell_width in zip(cells, values, widths):
            set_width(cell, cell_width)
            put_text(cell, value, size=7.5)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        fill(cells[2], COLORS[status])
        if tor_status in SUMMARY_COLORS:
            fill(cells[4], SUMMARY_COLORS[tor_status])

    output_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_docx)


def main() -> None:
    parser = argparse.ArgumentParser(description="Append TOR/SOC audit results to a copy of a DOCX")
    parser.add_argument("input_docx", type=Path)
    parser.add_argument("results_json", type=Path)
    parser.add_argument("output_docx", type=Path)
    args = parser.parse_args()
    build(args.input_docx, args.results_json, args.output_docx)
    print(args.output_docx)


if __name__ == "__main__":
    main()
