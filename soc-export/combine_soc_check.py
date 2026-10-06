"""Build a job's combined SOC_Check document (ADR 0008, ticket 10).

The web app writes the latest rows of every checked major item into one
results.json and calls this script. The results are appended to a copy of the
original SOC by the SOC skill's own append_results_to_docx.py, unchanged (not
AI). This script only adds a short note under the results heading: how many
major items are checked, where each one's rows came from, and which major
items are not checked yet.

    python combine_soc_check.py SKILL_SCRIPTS_DIR SOC.docx results.json manifest.json OUT.docx

manifest.json: {"checked": [{"label", "title", "rowCount", "checkedAt",
"model", "skillVersion"}], "unchecked": [{"label", "title"}]}
"""
from __future__ import annotations

import copy
import json
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph

# The skill's fixed subtitle under its results title; the note goes after it.
SUBTITLE_MARK = "ผลรายข้อแนบท้ายเอกสารต้นฉบับ"
FONT = "TH Sarabun New"


def item_name(item: dict) -> str:
    title = " ".join(str(item.get("title") or "").split())
    if len(title) > 60:
        title = title[:59].rstrip() + "…"
    return f"ข้อ {item['label']}" + (f" {title}" if title else "")


def note_lines(manifest: dict) -> list[tuple[str, bool, bool]]:
    """(text, bold, warn) per paragraph of the note."""
    checked, unchecked = manifest["checked"], manifest["unchecked"]
    total = len(checked) + len(unchecked)
    lines = [(f"ตรวจแล้ว {len(checked)}/{total} ข้อใหญ่ แต่ละข้อใหญ่แสดงเฉพาะผลตรวจครั้งล่าสุด", True, False)]
    for item in checked:
        source = " · ".join(str(v) for v in (item.get("checkedAt"), item.get("model"), item.get("skillVersion") and f"skill {item['skillVersion']}") if v)
        lines.append((f"{item_name(item)}: {item['rowCount']} แถว" + (f" ({source})" if source else ""), False, False))
    if unchecked:
        lines.append((f"ยังไม่ได้ตรวจ {len(unchecked)} ข้อใหญ่ (ไม่มีผลในเอกสารนี้): " + ", ".join(item_name(i) for i in unchecked), True, True))
    return lines


def add_note(output: Path, manifest: dict) -> None:
    doc = Document(output)
    subtitle = next((p for p in reversed(doc.paragraphs) if SUBTITLE_MARK in p.text), None)
    if subtitle is None:
        raise ValueError("results subtitle not found in the SOC_Check output")
    anchor = subtitle._p
    for text, bold, warn in note_lines(manifest):
        element = OxmlElement("w:p")
        if subtitle._p.pPr is not None:
            element.append(copy.deepcopy(subtitle._p.pPr))  # centred, like the subtitle
        anchor.addnext(element)
        anchor = element
        para = Paragraph(element, subtitle._parent)
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(0)
        run = para.add_run(text)
        run.bold = bold
        run.font.name = FONT
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
        run.font.size = Pt(10)
        if warn:
            run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    doc.save(output)


def main() -> None:
    if len(sys.argv) != 6:
        raise SystemExit(__doc__)
    skill_dir, soc, results, manifest, output = (Path(a) for a in sys.argv[1:])
    sys.path.insert(0, str(skill_dir.resolve()))
    from append_results_to_docx import build  # the skill's own script

    build(soc, results, output)
    add_note(output, json.loads(manifest.read_text(encoding="utf-8")))


if __name__ == "__main__":
    main()
